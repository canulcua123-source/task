#!/usr/bin/env python3
"""
Gestor de Tareas Escolares — UTM
Gestiona tareas por asignatura y genera documentos Word con portada institucional.
"""

import json
import os
import re
import sys
import tempfile
from datetime import datetime, date
from pathlib import Path

# Rutas base
BASE_DIR = Path(__file__).parent
CONFIG_PATH = BASE_DIR / "config.json"
DB_PATH = BASE_DIR / "tareas_db.json"
ENTREGAS_DIR = BASE_DIR / "sandbox" / "entregas"

ENTREGAS_DIR.mkdir(exist_ok=True)


def load_config():
    if not CONFIG_PATH.exists():
        print(f"Error: No se encontró {CONFIG_PATH}. Crea config.json primero.")
        sys.exit(1)
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def load_db():
    if not DB_PATH.exists():
        return {"tareas": []}
    with open(DB_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def _get_asignatura(clave):
    config = load_config()
    for a in config["asignaturas"]:
        if a["clave"] == clave:
            return a
    claves = [a["clave"] for a in config["asignaturas"]]
    print(f"Error: Asignatura '{clave}' no encontrada. Disponibles: {claves}")
    return None


def save_db(db):
    """Escritura atómica: escribe a temporal y reemplaza."""
    tmp_fd, tmp_path = tempfile.mkstemp(dir=BASE_DIR, suffix=".tmp.json")
    try:
        with open(tmp_fd, "w", encoding="utf-8") as f:
            json.dump(db, f, indent=2, ensure_ascii=False)
        os.replace(tmp_path, DB_PATH)
    except Exception:
        if Path(tmp_path).exists():
            Path(tmp_path).unlink()
        raise


# ─── Gestión de tareas ────────────────────────────────────────

def agregar_tarea(asignatura_clave, titulo, descripcion="", fecha_entrega=None, contenido=None, **extra):
    """Agrega una tarea nueva a la base de datos."""
    asignatura = _get_asignatura(asignatura_clave)
    if not asignatura:
        return None

    db = load_db()
    tarea = {
        "id": max((t["id"] for t in db["tareas"]), default=0) + 1,
        "asignatura_clave": asignatura_clave,
        "asignatura": asignatura["nombre"],
        "maestro": asignatura["maestro"],
        "titulo": titulo,
        "descripcion": descripcion,
        "contenido": contenido or [],
        "fecha_creacion": date.today().isoformat(),
        "fecha_entrega": fecha_entrega,
        "estado": "pendiente",
        "archivo_generado": None,
        **extra,
    }

    db["tareas"].append(tarea)
    save_db(db)
    print(f"Tarea #{tarea['id']} agregada: {titulo} ({asignatura['nombre']})")
    return tarea


def listar_tareas(filtro_estado=None, filtro_asignatura=None):
    """Lista todas las tareas, opcionalmente filtradas."""
    db = load_db()
    tareas = db["tareas"]

    if filtro_estado:
        tareas = [t for t in tareas if t["estado"] == filtro_estado]
    if filtro_asignatura:
        tareas = [t for t in tareas if t["asignatura_clave"] == filtro_asignatura]

    if not tareas:
        print("No hay tareas registradas con esos filtros.")
        return []

    print(f"\n{'ID':<4} {'Estado':<12} {'Asignatura':<20} {'Título':<30} {'Entrega':<12}")
    print("─" * 80)
    for t in tareas:
        estado_icon = {"pendiente": "⏳", "en_progreso": "🔧", "completada": "✅", "entregada": "📤"}.get(t["estado"], "?")
        entrega = t.get("fecha_entrega") or "Sin fecha"
        asig = t["asignatura"][:18]
        titulo = t["titulo"][:28]
        print(f"{t['id']:<4} {estado_icon} {t['estado']:<10} {asig:<20} {titulo:<30} {entrega:<12}")

    return tareas


ESTADOS_VALIDOS = {"pendiente", "en_progreso", "completada", "entregada"}

def actualizar_estado(tarea_id, nuevo_estado):
    """Cambia el estado de una tarea."""
    if nuevo_estado not in ESTADOS_VALIDOS:
        print(f"Error: Estado '{nuevo_estado}' no válido. Opciones: {ESTADOS_VALIDOS}")
        return None
    db = load_db()
    for t in db["tareas"]:
        if t["id"] == tarea_id:
            t["estado"] = nuevo_estado
            save_db(db)
            print(f"Tarea #{tarea_id} actualizada a: {nuevo_estado}")
            return t
    print(f"Tarea #{tarea_id} no encontrada.")
    return None


def obtener_tarea(tarea_id):
    """Obtiene una tarea por ID."""
    db = load_db()
    for t in db["tareas"]:
        if t["id"] == tarea_id:
            return t
    return None


def resumen_tareas():
    """Devuelve un resumen para uso agéntico."""
    db = load_db()
    config = load_config()
    tareas = db["tareas"]

    resumen = {
        "total": len(tareas),
        "pendientes": len([t for t in tareas if t["estado"] == "pendiente"]),
        "en_progreso": len([t for t in tareas if t["estado"] == "en_progreso"]),
        "completadas": len([t for t in tareas if t["estado"] == "completada"]),
        "entregadas": len([t for t in tareas if t["estado"] == "entregada"]),
        "por_asignatura": {},
        "proximas_entregas": []
    }

    for t in tareas:
        clave = t["asignatura_clave"]
        if clave not in resumen["por_asignatura"]:
            resumen["por_asignatura"][clave] = {"nombre": t["asignatura"], "pendientes": 0, "total": 0}
        resumen["por_asignatura"][clave]["total"] += 1
        if t["estado"] in ("pendiente", "en_progreso"):
            resumen["por_asignatura"][clave]["pendientes"] += 1

        if t.get("fecha_entrega") and t["estado"] in ("pendiente", "en_progreso"):
            resumen["proximas_entregas"].append({
                "id": t["id"],
                "titulo": t["titulo"],
                "asignatura": t["asignatura"],
                "fecha_entrega": t["fecha_entrega"],
                "estado": t["estado"]
            })

    resumen["proximas_entregas"].sort(key=lambda x: x["fecha_entrega"])
    return resumen


# ─── Generación de documento Word ─────────────────────────────

def _crear_portada_imagen(config, tarea):
    """Genera una imagen PNG de la portada estilo UTM con barras azules."""
    from PIL import Image, ImageDraw, ImageFont

    W, H = 2480, 3508  # A4 a 300 DPI
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)

    # ── Colores ──
    azul_oscuro = (0, 51, 102)
    azul_medio = (41, 82, 133)
    azul_claro = (90, 140, 190)
    azul_palido = (160, 200, 230)

    # ── Barras izquierda (anchas, con espacio entre ellas) ──
    # Barra 1: gruesa oscura
    draw.rectangle([100, 0, 260, H], fill=azul_oscuro)
    # Barra 2: media
    draw.rectangle([280, 0, 380, H], fill=azul_medio)
    # Barra 3: clara más delgada
    draw.rectangle([400, 0, 460, H], fill=azul_claro)
    # Barra 4: pálida delgada
    draw.rectangle([480, 0, 520, H], fill=azul_palido)

    # ── Barras esquina inferior derecha (espejo) ──
    bar_bottom_start = H - 600
    draw.rectangle([W - 520, bar_bottom_start + 200, W - 480, H], fill=azul_palido)
    draw.rectangle([W - 460, bar_bottom_start + 100, W - 400, H], fill=azul_claro)
    draw.rectangle([W - 380, bar_bottom_start + 50, W - 280, H], fill=azul_medio)
    draw.rectangle([W - 260, bar_bottom_start, W - 100, H], fill=azul_oscuro)

    # ── Fuentes ──
    def load_fonts():
        paths = [
            ("/Library/Fonts/Arial Bold.ttf", "/Library/Fonts/Arial.ttf"),
            ("/System/Library/Fonts/Supplemental/Arial Bold.ttf",
             "/System/Library/Fonts/Supplemental/Arial.ttf"),
        ]
        for bold_path, reg_path in paths:
            try:
                return {
                    "uni": ImageFont.truetype(bold_path, 100),
                    "carrera": ImageFont.truetype(bold_path, 68),
                    "campo_label": ImageFont.truetype(bold_path, 56),
                    "campo_valor": ImageFont.truetype(reg_path, 52),
                    "utm": ImageFont.truetype(bold_path, 140),
                }
            except OSError:
                continue
        # Fallback Helvetica
        try:
            return {
                "uni": ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 100),
                "carrera": ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 68),
                "campo_label": ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 56),
                "campo_valor": ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 52),
                "utm": ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 140),
            }
        except OSError:
            f = ImageFont.load_default()
            return {"uni": f, "carrera": f, "campo_label": f, "campo_valor": f, "utm": f}

    fonts = load_fonts()
    est = config["estudiante"]
    esc = config["escuela"]

    # Área de contenido (a la derecha de las barras)
    cl = 580  # content left
    cr = W - 200  # content right
    cc = (cl + cr) // 2  # content center
    cw = cr - cl  # content width

    def draw_centered(text, y, font, color=azul_oscuro):
        bbox = draw.textbbox((0, 0), text, font=font)
        tw = bbox[2] - bbox[0]
        draw.text((cc - tw // 2, y), text, fill=color, font=font)

    # ── Construir bloques para calcular altura total ──
    # Cada bloque: (tipo, datos, altura)
    blocks = []

    # Logo / UTM
    logo_path = BASE_DIR / esc.get("logo", "")
    has_logo = logo_path.exists() and logo_path.is_file()
    if has_logo:
        blocks.append(("logo", logo_path, 460))
    else:
        blocks.append(("utm_text", None, 200))

    # Universidad
    uni_lines = _wrap_text_center(draw, esc["nombre"].upper(), fonts["uni"], cw)
    blocks.append(("uni", uni_lines, len(uni_lines) * 120))

    # Carrera
    car_lines = _wrap_text_center(draw, est["carrera"].upper(), fonts["carrera"], cw)
    blocks.append(("carrera", car_lines, len(car_lines) * 85))

    # Campos
    campos = [
        ("Actividad:", tarea["titulo"]),
        ("Materia:", tarea["asignatura"]),
        ("Maestro:", tarea["maestro"]),
        ("Integrantes:", est["nombre"]),
    ]
    for etiqueta, valor in campos:
        full = f"{etiqueta} {valor}"
        full_lines = _wrap_text_center(draw, full, fonts["campo_label"], cw - 80)
        if len(full_lines) == 1:
            blocks.append(("campo_inline", (etiqueta, valor), 70))
        else:
            val_lines = _wrap_text_center(draw, valor, fonts["campo_valor"], cw - 100)
            blocks.append(("campo_multi", (etiqueta, val_lines), 70 + len(val_lines) * 65))

    # Grado + Grupo
    blocks.append(("grado_grupo", None, 160))

    # Fecha
    fecha_str = None
    if tarea.get("fecha_entrega"):
        try:
            fecha = datetime.strptime(tarea["fecha_entrega"], "%Y-%m-%d")
            meses = ["enero","febrero","marzo","abril","mayo","junio",
                     "julio","agosto","septiembre","octubre","noviembre","diciembre"]
            fecha_str = f"Fecha de entrega: {fecha.day} {meses[fecha.month-1]} {fecha.year}"
        except ValueError:
            fecha_str = f"Fecha de entrega: {tarea['fecha_entrega']}"
        blocks.append(("fecha", fecha_str, 70))

    # Calcular altura total y distribuir espaciado
    total_content_h = sum(b[2] for b in blocks)
    usable_h = H - 300  # margen arriba y abajo de 150px
    n_gaps = len(blocks) - 1
    if total_content_h < usable_h and n_gaps > 0:
        gap = (usable_h - total_content_h) // n_gaps
        gap = min(gap, 180)  # máximo espacio entre bloques
    else:
        gap = 40

    # Recalcular inicio para centrar verticalmente
    total_with_gaps = total_content_h + gap * n_gaps
    y = max(150, (H - total_with_gaps) // 2)

    # ── Dibujar bloques ──
    for block in blocks:
        tipo = block[0]

        if tipo == "logo":
            try:
                logo = Image.open(str(block[1]))
                logo_max_h = 400
                ratio = logo_max_h / logo.height
                logo_w = int(logo.width * ratio)
                logo = logo.resize((logo_w, logo_max_h), Image.LANCZOS)
                logo_x = cc - logo_w // 2
                if logo.mode == "RGBA":
                    img.paste(logo, (logo_x, y), logo)
                else:
                    img.paste(logo, (logo_x, y))
            except Exception:
                pass
            y += block[2] + gap

        elif tipo == "utm_text":
            draw_centered("UTM", y, fonts["utm"], color=(0, 102, 51))
            y += block[2] + gap

        elif tipo == "uni":
            for line in block[1]:
                draw_centered(line, y, fonts["uni"])
                y += 120
            y += gap

        elif tipo == "carrera":
            for line in block[1]:
                draw_centered(line, y, fonts["carrera"])
                y += 85
            y += gap

        elif tipo == "campo_inline":
            etiqueta, valor = block[1]
            etiq_bbox = draw.textbbox((0, 0), etiqueta + " ", font=fonts["campo_label"])
            val_bbox = draw.textbbox((0, 0), valor, font=fonts["campo_valor"])
            total_w = (etiq_bbox[2] - etiq_bbox[0]) + (val_bbox[2] - val_bbox[0])
            x_start = cc - total_w // 2
            draw.text((x_start, y), etiqueta + " ", fill=azul_oscuro, font=fonts["campo_label"])
            draw.text((x_start + etiq_bbox[2] - etiq_bbox[0], y), valor, fill=azul_oscuro, font=fonts["campo_valor"])
            y += block[2] + gap

        elif tipo == "campo_multi":
            etiqueta, val_lines = block[1]
            draw_centered(etiqueta, y, fonts["campo_label"])
            y += 70
            for vl in val_lines:
                draw_centered(vl, y, fonts["campo_valor"])
                y += 65
            y += gap

        elif tipo == "grado_grupo":
            draw_centered(f"Grado: {est['cuatrimestre']}", y, fonts["campo_label"])
            y += 80
            draw_centered(f"Grupo: {est['grupo']}", y, fonts["campo_label"])
            y += 80 + gap

        elif tipo == "fecha":
            draw_centered(block[1], y, fonts["campo_label"])
            y += block[2] + gap

    # Guardar imagen temporal (tempfile seguro, sin race condition)
    fd, portada_path = tempfile.mkstemp(suffix=".png", prefix="portada_", dir=BASE_DIR / "assets")
    os.close(fd)
    img.save(portada_path, "PNG", quality=95)
    return portada_path


def _wrap_text_center(draw, text, font, max_width):
    """Divide texto en líneas que quepan en max_width."""
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        bbox = draw.textbbox((0, 0), test, font=font)
        if bbox[2] - bbox[0] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def generar_documento(tarea_id):
    """Genera un documento Word con portada estilo UTM para una tarea."""
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor, Emu
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    tarea = obtener_tarea(tarea_id)
    if not tarea:
        print(f"Tarea #{tarea_id} no encontrada.")
        return None

    config = load_config()
    est = config["estudiante"]
    esc = config["escuela"]
    fmt = config["formato"]

    doc = Document()

    # Configurar márgenes para portada (mínimos, la imagen llena la página)
    section = doc.sections[0]
    section.top_margin = Cm(0)
    section.bottom_margin = Cm(0)
    section.left_margin = Cm(0)
    section.right_margin = Cm(0)
    page_w = section.page_width
    page_h = section.page_height

    color_azul = RGBColor(0, 0x33, 0x66)

    # ── PORTADA como imagen ──
    portada_path = _crear_portada_imagen(config, tarea)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(portada_path, width=page_w, height=page_h)

    # ── CONTENIDO (con márgenes normales) ──
    contenido = tarea.get("contenido")
    if isinstance(contenido, list) and contenido:
        # Nueva sección ya hace salto de página, no agregar page_break extra
        new_section = doc.add_section()
        new_section.top_margin = Cm(2.5)
        new_section.bottom_margin = Cm(2.5)
        new_section.left_margin = Cm(2.5)
        new_section.right_margin = Cm(2.5)

        for seccion in contenido:
            tipo = seccion.get("tipo", "parrafo")

            if tipo == "titulo":
                p = doc.add_heading(seccion["texto"], level=1)
                for run in p.runs:
                    run.font.color.rgb = color_azul
                    run.font.name = fmt["fuente"]

            elif tipo == "subtitulo":
                p = doc.add_heading(seccion["texto"], level=2)
                for run in p.runs:
                    run.font.color.rgb = color_azul
                    run.font.name = fmt["fuente"]

            elif tipo == "parrafo":
                p = doc.add_paragraph()
                run = p.add_run(seccion["texto"])
                run.font.size = Pt(fmt["tamano_cuerpo"])
                run.font.name = fmt["fuente"]
                p.paragraph_format.line_spacing = 1.5
                p.paragraph_format.first_line_indent = Cm(1.25)

            elif tipo == "lista":
                for item in seccion.get("items", []):
                    p = doc.add_paragraph(item, style="List Bullet")
                    for run in p.runs:
                        run.font.size = Pt(fmt["tamano_cuerpo"])
                        run.font.name = fmt["fuente"]

    # Guardar
    titulo_safe = re.sub(r'[^\w\-]', '_', tarea['titulo'])[:30]
    nombre_archivo = f"{tarea['asignatura_clave']}_{tarea['id']}_{titulo_safe}.docx"
    ruta = ENTREGAS_DIR / nombre_archivo
    doc.save(str(ruta))

    # Limpiar temporal
    if Path(portada_path).exists():
        Path(portada_path).unlink()

    # Actualizar DB
    db = load_db()
    for t in db["tareas"]:
        if t["id"] == tarea_id:
            t["archivo_generado"] = str(ruta)
            break
    save_db(db)

    print(f"Documento generado: {ruta}")
    return str(ruta)


# ─── Integración Teams / utilidades ──────────────────────────

def importar_desde_teams(teams_data, asignatura_clave):
    titulo = teams_data.get("title", "Sin título")
    descripcion = teams_data.get("description", "")
    instrucciones = teams_data.get("instructions", "")

    fecha_entrega = None
    if teams_data.get("dueDateTime"):
        try:
            dt = datetime.fromisoformat(teams_data["dueDateTime"].replace("Z", "+00:00"))
            fecha_entrega = dt.date().isoformat()
        except (ValueError, AttributeError):
            fecha_entrega = teams_data["dueDateTime"]

    contenido = []
    if instrucciones:
        contenido.append({"tipo": "titulo", "texto": titulo})
        contenido.append({"tipo": "parrafo", "texto": instrucciones})

    return agregar_tarea(
        asignatura_clave, titulo, descripcion, fecha_entrega, contenido,
        teams_data=teams_data,
        class_id=teams_data.get("class_id"),
        assignment_id=teams_data.get("assignment_id"),
    )


def generar_pendientes():
    db = load_db()
    generados = []
    for t in db["tareas"]:
        if (
            t["estado"] != "entregada"
            and isinstance(t.get("contenido"), list)
            and t["contenido"]
            and not t.get("archivo_generado")
        ):
            ruta = generar_documento(t["id"])
            if ruta:
                generados.append(ruta)
    return generados


def alertas_vencidas():
    db = load_db()
    hoy = date.today()
    alertas = []
    for t in db["tareas"]:
        if not t.get("fecha_entrega") or t["estado"] in ("completada", "entregada"):
            continue
        try:
            fecha = date.fromisoformat(t["fecha_entrega"])
        except ValueError:
            continue
        dias = (fecha - hoy).days
        if dias <= 3:
            alertas.append({
                "id": t["id"],
                "titulo": t["titulo"],
                "asignatura": t["asignatura"],
                "fecha_entrega": t["fecha_entrega"],
                "estado": t["estado"],
                "dias_restantes": dias,
            })
    alertas.sort(key=lambda x: x["dias_restantes"])
    return alertas


def flujo_completo(tarea_id, subir_teams=False):
    """Orquesta: generar doc → actualizar estado → opcionalmente subir a Teams."""
    tarea = obtener_tarea(tarea_id)
    if not tarea:
        print(f"Tarea #{tarea_id} no encontrada.")
        return None

    # Verificar que tenga contenido
    if not isinstance(tarea.get("contenido"), list) or not tarea["contenido"]:
        print(f"Tarea #{tarea_id} no tiene contenido para generar.")
        return None

    # Generar documento
    ruta = generar_documento(tarea_id)
    if not ruta:
        return None

    # Actualizar estado
    actualizar_estado(tarea_id, "completada")

    # Subir a Teams si se pide y tiene datos de Teams
    if subir_teams and tarea.get("class_id") and tarea.get("assignment_id"):
        try:
            from teams import subir_archivo_tarea, entregar_tarea
            print(f"Subiendo a Teams...")
            subir_archivo_tarea(tarea["class_id"], tarea["assignment_id"], ruta)
            entregar_tarea(tarea["class_id"], tarea["assignment_id"])
            actualizar_estado(tarea_id, "entregada")
            print(f"Tarea #{tarea_id} entregada en Teams.")
        except Exception as e:
            print(f"Error al subir a Teams: {e}")

    return ruta


# ─── CLI ──────────────────────────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("""
Gestor de Tareas UTM
Uso:
  python gestor.py listar [pendiente|completada|entregada] [asignatura_clave]
  python gestor.py resumen
  python gestor.py agregar <clave_asignatura> <titulo> [fecha_entrega YYYY-MM-DD]
  python gestor.py estado <id> <nuevo_estado>
  python gestor.py generar <id>
  python gestor.py asignaturas
        """)
        return

    cmd = sys.argv[1]

    if cmd == "listar":
        estado = sys.argv[2] if len(sys.argv) > 2 else None
        asig = sys.argv[3] if len(sys.argv) > 3 else None
        listar_tareas(estado, asig)

    elif cmd == "resumen":
        r = resumen_tareas()
        print(f"\nTotal: {r['total']} | Pendientes: {r['pendientes']} | En progreso: {r['en_progreso']} | Completadas: {r['completadas']} | Entregadas: {r['entregadas']}")
        if r["proximas_entregas"]:
            print("\nPróximas entregas:")
            for e in r["proximas_entregas"]:
                print(f"  #{e['id']} {e['titulo']} ({e['asignatura']}) — {e['fecha_entrega']}")

    elif cmd == "agregar":
        if len(sys.argv) < 4:
            print("Uso: python gestor.py agregar <clave_asignatura> <titulo> [fecha_entrega]")
            return
        fecha = sys.argv[4] if len(sys.argv) > 4 else None
        agregar_tarea(sys.argv[2], sys.argv[3], fecha_entrega=fecha)

    elif cmd == "estado":
        if len(sys.argv) < 4:
            print("Uso: python gestor.py estado <id> <nuevo_estado>")
            return
        actualizar_estado(int(sys.argv[2]), sys.argv[3])

    elif cmd == "generar":
        if len(sys.argv) < 3:
            print("Uso: python gestor.py generar <id>")
            return
        generar_documento(int(sys.argv[2]))

    elif cmd == "asignaturas":
        config = load_config()
        print("\nAsignaturas registradas:")
        for a in config["asignaturas"]:
            print(f"  [{a['clave']}] {a['nombre']} — {a['maestro']}")

    elif cmd == "pendientes":
        rutas = generar_pendientes()
        if rutas:
            print(f"\n{len(rutas)} documento(s) generado(s):")
            for r in rutas:
                print(f"  {r}")
        else:
            print("No hay tareas pendientes con contenido listo para generar.")

    elif cmd == "alertas":
        alertas = alertas_vencidas()
        if not alertas:
            print("No hay tareas vencidas ni próximas a vencer.")
        else:
            print(f"\n{'ID':<4} {'Días':<6} {'Entrega':<12} {'Asignatura':<20} Título")
            print("─" * 70)
            for a in alertas:
                dias = a["dias_restantes"]
                if dias < 0:
                    label = f"VENCIDA ({abs(dias)}d)"
                elif dias == 0:
                    label = "HOY"
                else:
                    label = f"{dias}d"
                print(f"{a['id']:<4} {label:<6} {a['fecha_entrega']:<12} {a['asignatura'][:18]:<20} {a['titulo']}")

    elif cmd == "flujo":
        if len(sys.argv) < 3:
            print("Uso: python gestor.py flujo <id> [--teams]")
            return
        subir = "--teams" in sys.argv
        ruta = flujo_completo(int(sys.argv[2]), subir_teams=subir)
        if ruta:
            print(f"Flujo completo: {ruta}")


if __name__ == "__main__":
    main()
